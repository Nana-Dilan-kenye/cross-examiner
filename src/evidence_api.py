from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
import csv
import json

ROOT = Path(__file__).resolve().parent.parent
CSV_DATA = ROOT / "data" / "coverage_evidence.csv"
SURVEY_DATA = ROOT / "data" / "survey_figures.csv"
XLSX_DATA = ROOT / "data" / "wuenic2025rev_web-update.xlsx"
WUENIC_URL = "https://data.unicef.org/resources/dataset/immunization/"

# The countries, antigen and years the workbook extract is asked for.
WORKBOOK_COUNTRIES = ("CMR", "ETH")
WORKBOOK_ANTIGEN = "mcv1"
WORKBOOK_TARGET_GROUP = "si"
WORKBOOK_YEARS = (2017, 2018, 2020, 2021, 2023)

# The workbook is large and slow to parse, so the records are cached after
# the first read. Without this every request spends ~25s in openpyxl.
_CSV_CACHE = None
_WUENIC_CACHE = None
_SURVEY_CACHE = None


def load_csv_evidence():
    global _CSV_CACHE
    if _CSV_CACHE is not None:
        return _CSV_CACHE
    with CSV_DATA.open(newline="", encoding="utf-8") as f:
        _CSV_CACHE = list(csv.DictReader(f))
    return _CSV_CACHE


def load_wuenic():
    global _WUENIC_CACHE
    if _WUENIC_CACHE is not None:
        return _WUENIC_CACHE

    if not XLSX_DATA.exists():
        _WUENIC_CACHE = []
        return _WUENIC_CACHE
    wb = load_workbook(XLSX_DATA, read_only=True, data_only=True)
    try:
        ws = wb["country"]
        rows = ws.iter_rows(values_only=True)
        next(rows)  # Header

        records = []
        for row in rows:
            if (
                row[0] in WORKBOOK_COUNTRIES
                and row[2] == WORKBOOK_ANTIGEN
                and row[4] == WORKBOOK_TARGET_GROUP
                and row[3] in WORKBOOK_YEARS
                and row[5] is not None
            ):
                records.append({
                    "country": row[1],
                    "antigen": "MCV1",
                    "year": str(row[3]),
                    "measure": "WUENIC estimate",
                    "value_percent": str(row[5]),
                    "source": str(row[6]),
                    "source_url": WUENIC_URL,
                    "notes": f"Workbook target group: {row[4]}",
                })
        _WUENIC_CACHE = records
        return records
    finally:
        wb.close()


def load_survey_figures():
    """The compiled sheet of published household-survey figures.

    These are transcribed by hand from the survey reports, because the
    reports are not machine-readable in this environment. The sheet is
    served in the same shape as the other endpoints so the unified layer
    can read it as one more source.
    """
    global _SURVEY_CACHE
    if _SURVEY_CACHE is not None:
        return _SURVEY_CACHE
    records = []
    if SURVEY_DATA.exists():
        with SURVEY_DATA.open(newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                records.append({
                    "country": row.get("country", ""),
                    "antigen": row.get("antigen", ""),
                    "year": row.get("year", ""),
                    "measure": row.get("measure", ""),
                    "value_percent": row.get("value_percent", ""),
                    "source": row.get("source_id", ""),
                    "source_url": row.get("source_ref", ""),
                    "notes": row.get("notes", ""),
                })
    _SURVEY_CACHE = records
    return records


class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        route = urlparse(self.path).path

        if route == "/evidence":
            records = load_csv_evidence()
        elif route == "/wuenic":
            records = load_wuenic()
        elif route == "/survey_figures":
            records = load_survey_figures()
        else:
            self.send_error(404)
            return

        body = json.dumps({"evidence": records}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        print("%s - %s" % (self.address_string(), format % args))


if __name__ == "__main__":
    print("Cross-Examiner evidence API listening on port 8765")
    HTTPServer(("0.0.0.0", 8765), Handler).serve_forever()
